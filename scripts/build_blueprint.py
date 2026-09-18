from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Oracle_Patching_Utility_Architecture_and_Delivery_Blueprint.docx"
ASSET_DIR = ROOT / ".document_assets"
ASSET_DIR.mkdir(exist_ok=True)

NAVY = "203748"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2933"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHTER_BLUE = "F3F7FB"
MID_GRAY = "D0D5DD"
GOLD = "B5852B"
RED = "9B1C1C"
GREEN = "22643A"

USABLE_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM_DXA = 80
CELL_SIDE_DXA = 120


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=MID_GRAY, size="4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != USABLE_DXA:
        raise ValueError(f"table width must total {USABLE_DXA}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell, CELL_TOP_BOTTOM_DXA, CELL_SIDE_DXA, CELL_TOP_BOTTOM_DXA, CELL_SIDE_DXA)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=None, color=INK, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def paragraph_border_bottom(paragraph, color=BLUE, size="14", space="5") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(paragraph, instruction: str, display="1") -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_custom_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    existing_abs = [int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs or [0]) + 1
    next_num = max(existing_num or [0]) + 1

    def build(kind: str, abs_id: int, num_id: int) -> int:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        nsid = OxmlElement("w:nsid")
        nsid.set(qn("w:val"), f"{abs_id:08X}")
        abstract.append(nsid)
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
        lvl.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        lvl.append(p_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abs_id))
        num.append(ref)
        numbering.append(num)
        return num_id

    bullet_id = build("bullet", next_abs, next_num)
    number_id = build("decimal", next_abs + 1, next_num + 1)
    return bullet_id, number_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def configure_styles(doc: Document) -> tuple[int, int]:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True
    return add_custom_numbering(doc)


def set_headers_footers(doc: Document) -> None:
    for section in doc.sections:
        header_p = section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_p.paragraph_format.space_after = Pt(2)
        r = header_p.add_run("ORACLE PATCHING UTILITY  |  ARCHITECTURE & DELIVERY BLUEPRINT")
        set_run_font(r, size=8.5, color=MUTED, bold=True)
        paragraph_border_bottom(header_p, color=MID_GRAY, size="4", space="4")

        footer_p = section.footer.paragraphs[0]
        footer_p.paragraph_format.space_before = Pt(2)
        footer_p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        left = footer_p.add_run("Version 0.1  |  9 July 2026")
        set_run_font(left, size=8.5, color=MUTED)
        footer_p.add_run("\t")
        page_label = footer_p.add_run("Page ")
        set_run_font(page_label, size=8.5, color=MUTED)
        add_field(footer_p, "PAGE")


def add_title(doc, text: str, size=30, color=NAVY, after=8, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=True)
    return p


def add_subtitle(doc, text: str, size=15, color=DARK_BLUE, after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=size, color=color)
    return p


def add_body(doc, text: str, bold_lead: str | None = None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_bullet(doc, text: str, bullet_num_id: int, bold_lead: str | None = None):
    p = doc.add_paragraph()
    apply_numbering(p, bullet_num_id)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_callout(doc, label: str, text: str, fill=LIGHTER_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [USABLE_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=accent, size="8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10.5, color=accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers: list[str], rows: list[list[str]], widths_dxa: list[int],
              font_size=9.5, header_fill=LIGHT_BLUE, alternate=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_cant_split(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(text)
        set_run_font(r, size=font_size, color=NAVY, bold=True)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        set_cant_split(table.rows[-1])
        for c_idx, value in enumerate(row):
            cell = cells[c_idx]
            if alternate and r_idx % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            r = p.add_run(value)
            set_run_font(r, size=font_size, color=INK, bold=(c_idx == 0 and len(value) < 16))
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def set_picture_alt_text(shape, title: str, description: str) -> None:
    inline = shape._inline
    doc_pr = inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


FONT_REGULAR = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


def pil_font(size: int, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size=size)


def centered_multiline(draw, xy, text, font, fill, spacing=8):
    box = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    width = box[2] - box[0]
    height = box[3] - box[1]
    draw.multiline_text((xy[0] - width / 2, xy[1] - height / 2), text, font=font,
                        fill=fill, spacing=spacing, align="center")


def draw_arrow(draw, start, end, color="#667085", width=6, label=""):
    x1, y1 = start
    x2, y2 = end
    draw.line([start, end], fill=color, width=width)
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    tip = (x2, y2)
    back = (x2 - ux * 28, y2 - uy * 28)
    left = (back[0] + px * 15, back[1] + py * 15)
    right = (back[0] - px * 15, back[1] - py * 15)
    draw.polygon([tip, left, right], fill=color)
    if label:
        mid = ((x1 + x2) / 2, (y1 + y2) / 2 - 18)
        f = pil_font(22)
        bbox = draw.textbbox((0, 0), label, font=f)
        pad = 8
        draw.rounded_rectangle((mid[0] - (bbox[2] - bbox[0]) / 2 - pad,
                                mid[1] - (bbox[3] - bbox[1]) / 2 - pad,
                                mid[0] + (bbox[2] - bbox[0]) / 2 + pad,
                                mid[1] + (bbox[3] - bbox[1]) / 2 + pad),
                               radius=6, fill="white")
        draw.text((mid[0] - (bbox[2] - bbox[0]) / 2, mid[1] - (bbox[3] - bbox[1]) / 2),
                  label, font=f, fill=color)


def draw_architecture(path: Path) -> None:
    image = Image.new("RGB", (2160, 1210), "white")
    draw = ImageDraw.Draw(image)

    def box(x, y, w, h, title, subtitle="", face="#F3F7FB", edge="#2E74B5"):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=24, fill=face, outline=edge, width=5)
        centered_multiline(draw, (x + w / 2, y + h * 0.38), title, pil_font(34, True), "#203748")
        if subtitle:
            centered_multiline(draw, (x + w / 2, y + h * 0.72), subtitle, pil_font(25), "#4B5563", spacing=7)

    centered_multiline(draw, (1080, 80), "Deployment-neutral control plane with outbound managed agents",
                       pil_font(47, True), "#203748")
    box(60, 270, 390, 205, "Operators & Approvers", "Web UI / CLI\nEnterprise identity", "#FFF8E8", "#B5852B")
    box(575, 220, 1010, 295, "Control Plane", "API • Inventory • Reconciliation • Compliance\nPolicy • Workflow • Approvals • Audit", "#E8EEF5", "#2E74B5")
    box(1710, 270, 390, 205, "Artifact Store", "Local disk / NFS\nOptional S3 adapter", "#F4F6F9", "#667085")
    box(120, 710, 530, 245, "Managed Host Agent", "Dynamic discovery • Task leasing\nTyped operations • Evidence", "#ECF8F0", "#22643A")
    box(815, 710, 530, 245, "Oracle Estate", "Homes • Databases • Listeners\nData Guard • GI / RAC / ASM", "#FFF4F2", "#9B1C1C")
    box(1510, 710, 530, 245, "System of Record", "PostgreSQL\nSnapshots • Workflow • Audit", "#F4F6F9", "#667085")
    draw_arrow(draw, (450, 370), (575, 370))
    draw_arrow(draw, (1585, 370), (1710, 370))
    draw_arrow(draw, (800, 515), (500, 710), label="mTLS tasks & evidence")
    draw_arrow(draw, (650, 830), (815, 830), label="typed operations")
    draw_arrow(draw, (1345, 830), (1510, 830), label="observations")
    draw_arrow(draw, (1650, 710), (1400, 515), label="state & audit")
    centered_multiline(draw, (1080, 1115), "No inbound SSH • No arbitrary shell • Cloud services are optional adapters",
                       pil_font(31, True), "#2E74B5")
    image.save(path, optimize=True)


def draw_discovery_loop(path: Path) -> None:
    image = Image.new("RGB", (2160, 870), "white")
    draw = ImageDraw.Draw(image)
    titles = [
        ("1", "Collect", "OS + Oracle tools\nfixed read-only scripts"),
        ("2", "Snapshot", "immutable observations\nsource + time + digest"),
        ("3", "Reconcile", "stable identity\nconflict quarantine"),
        ("4", "Project", "current inventory\ntopology + history"),
        ("5", "Decide", "compliance + readiness\npolicy-gated workflow"),
    ]
    xs = [55, 485, 915, 1345, 1775]
    centered_multiline(draw, (1080, 72), "Evidence-backed discovery and reconciliation loop",
                       pil_font(46, True), "#203748")
    for idx, ((num, title, subtitle), x) in enumerate(zip(titles, xs)):
        draw.ellipse((x + 8, 185, x + 72, 249), fill="#2E74B5")
        centered_multiline(draw, (x + 40, 217), num, pil_font(29, True), "white")
        draw.rounded_rectangle((x, 285, x + 330, 610), radius=24, fill="#F3F7FB", outline="#2E74B5", width=5)
        centered_multiline(draw, (x + 165, 385), title, pil_font(37, True), "#203748")
        centered_multiline(draw, (x + 165, 500), subtitle, pil_font(26), "#4B5563", spacing=9)
        if idx < len(xs) - 1:
            draw_arrow(draw, (x + 330, 445), (xs[idx + 1] - 20, 445), width=6)
    centered_multiline(draw, (1080, 775), "A partial scan may reduce confidence; it never silently deletes a resource.",
                       pil_font(31, True), "#9B1C1C")
    image.save(path, optimize=True)


def parse_work_packages(path: Path):
    tracks = []
    current = None
    for raw in path.read_text(encoding="utf-8").splitlines():
        if raw.startswith("## Track "):
            current = {"title": raw[3:].strip(), "rows": []}
            tracks.append(current)
            continue
        if current and re.match(r"^\| [A-Z]+-\d+ \|", raw):
            parts = [p.strip() for p in raw.strip().strip("|").split("|")]
            if len(parts) == 4:
                current["rows"].append(parts)
    return tracks


def build_document() -> None:
    doc = Document()
    bullet_id, number_id = configure_styles(doc)
    set_headers_footers(doc)

    arch_img = ASSET_DIR / "architecture.png"
    discovery_img = ASSET_DIR / "discovery_loop.png"
    draw_architecture(arch_img)
    draw_discovery_loop(discovery_img)

    # Cover page - editorial_cover pattern with a restrained technical palette.
    for _ in range(4):
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    r = kicker.add_run("GREENFIELD PROGRAM BLUEPRINT")
    set_run_font(r, size=10.5, color=GOLD, bold=True)
    add_title(doc, "Oracle Patching Utility", size=30, color=NAVY, after=8)
    add_subtitle(doc, "Architecture, Dynamic Discovery, Safety Model and Delivery Plan", size=15, color=DARK_BLUE, after=28)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(64)
    r = p.add_run("An OEM-style managed-agent platform built independently for accurate, auditable Oracle patching")
    set_run_font(r, size=11, color=MUTED, italic=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Version 0.1")
    set_run_font(r, size=12, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("9 July 2026  |  Architecture baseline for implementation")
    set_run_font(r, size=10, color=MUTED)
    doc.add_page_break()

    # Document control and use.
    doc.add_heading("Document purpose", level=1)
    add_body(doc, "This document is the implementation baseline for a greenfield Oracle patching utility. It consolidates the agreed architecture, the OEM-style dynamic discovery approach, execution safeguards, major risks, accuracy gates, delivery waves, and all independently assignable work packages.")
    add_callout(doc, "Decision", "Begin with canonical identities and versioned contracts, then build the read-only agent and reconciliation pipeline. Real patch execution remains disabled until the simulator and Oracle lab gates pass.")
    add_table(doc,
              ["Document attribute", "Value"],
              [
                  ["Status", "Architecture baseline - ready to begin contract work"],
                  ["Initial target", "Oracle Database 19c on Oracle Linux x86-64; standalone first"],
                  ["Deployment", "On-premises capable; cloud services are optional adapters"],
                  ["Storage", "Local filesystem or NFS by default; provider-neutral ArtifactStore"],
                  ["Delivery model", "Modular control plane plus a separately installed managed host agent"],
                  ["First package", "CON-01 - Canonical resource identities"],
              ],
              [2200, 7160], font_size=9.5)

    doc.add_heading("How to use this blueprint", level=2)
    for item in [
        "Treat the program charter and safety rules as non-negotiable architecture constraints.",
        "Assign work by package ID from Appendix A; every package has declared dependencies and acceptance evidence.",
        "Freeze shared contracts before parallel teams implement agent, discovery, security, storage, and workflow components.",
        "Advance between delivery waves only after the corresponding accuracy and safety gate passes.",
        "Record decisions that change scope, identity, privileges, topology support, or recovery behavior as architecture decisions.",
    ]:
        add_bullet(doc, item, bullet_id)

    doc.add_page_break()
    doc.add_heading("Contents", level=1)
    contents = [
        "1. Executive summary", "2. Vision, scope and non-goals", "3. System architecture",
        "4. OEM-style dynamic discovery", "5. Inventory identity and reconciliation",
        "6. Patch lifecycle and execution safety", "7. Security, trust and deployment",
        "8. Principal challenges and mitigations", "9. Accuracy and quality gates",
        "10. Workstreams and parallel delivery waves", "11. First implementation package",
        "12. Governance and definition of done", "13. Decisions to validate",
        "Appendix A. Complete work-package register", "Appendix B. Core terminology",
    ]
    for item in contents:
        add_bullet(doc, item, bullet_id)
    doc.add_page_break()

    # Main narrative.
    doc.add_heading("1. Executive summary", level=1)
    add_body(doc, "The utility will use a dedicated host agent, similar in operating model to enterprise monitoring platforms, to discover Oracle installations dynamically and report evidence-backed observations to a central control plane. The platform will then reconcile those observations into a current inventory, compare that inventory with approved patch baselines, perform read-only readiness checks, and execute only explicitly approved patch workflows.")
    add_body(doc, "The agent is not a remote shell. It exposes only fixed, versioned operations with typed parameters. Discovery, prechecks and patch execution remain separate capabilities with separate privilege profiles. A control-plane timeout or lost heartbeat is never interpreted as an Oracle failure or as permission to retry a destructive operation.")
    add_body(doc, "Delivery is divided into 114 independently testable work packages. Parallel development begins only after canonical identifiers, resource schemas, result envelopes and protocol compatibility rules are agreed. This contract-first approach is the fastest path that does not sacrifice accuracy.")
    add_callout(doc, "Core outcome", "Discover accurately, decide transparently, execute deterministically, verify independently, and retain evidence for every transition.", fill="ECF8F0", accent=GREEN)

    doc.add_heading("2. Vision, scope and non-goals", level=1)
    doc.add_heading("2.1 Vision", level=2)
    add_body(doc, "Provide one trustworthy view of the Oracle estate and one controlled path from patch requirement to validated outcome. Administrators define desired policy and business metadata; the agent observes technical state dynamically. The platform never confuses desired state with observed fact.")
    doc.add_heading("2.2 Initial scope", level=2)
    for item in [
        "Oracle Database 19c on Oracle Linux x86-64.",
        "Standalone databases before Data Guard, then Grid Infrastructure and RAC.",
        "Dynamic discovery of hosts, Oracle homes, databases, CDB/PDBs, listeners, services, patch inventory and topology evidence.",
        "Patch catalog, approved baselines, compliance, readiness, approvals, scheduling, execution, validation, rollback coordination, reporting and audit.",
        "Local filesystem or NFS artifact storage for fully on-premises deployments; optional adapters may be added without changing domain records.",
    ]:
        add_bullet(doc, item, bullet_id)
    doc.add_heading("2.3 Explicit non-goals", level=2)
    for item in [
        "Generic SSH, arbitrary shell, uploaded scripts or user-supplied SQL.",
        "Automatic patching merely because a target is non-compliant.",
        "Assuming that every Oracle failure can be rolled back automatically.",
        "Reimplementing Oracle conflict or applicability logic when Oracle tooling is authoritative.",
        "Requiring Kubernetes, public cloud services or public runtime connectivity.",
        "AI-driven approval or command execution.",
    ]:
        add_bullet(doc, item, bullet_id)

    doc.add_heading("3. System architecture", level=1)
    add_body(doc, "The recommended starting architecture is a modular control plane and a separately installed host agent. API, worker and scheduler processes may be deployed separately, but initially share one codebase and PostgreSQL system of record. This limits distributed transaction risk while contracts are still evolving.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(arch_img), width=Inches(6.42))
    set_picture_alt_text(shape, "Oracle patching utility system architecture", "Operators use a control plane connected to provider-neutral artifact storage, PostgreSQL and outbound managed host agents that discover and operate on the Oracle estate.")
    cap = doc.add_paragraph("Figure 1. Deployment-neutral control plane and managed-agent boundary", style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    components = [
        ["Control Plane API", "Versioned APIs for inventory, baselines, changes, workflows, approvals, evidence and reports."],
        ["Identity and Policy", "User and workload identity, RBAC, separation of duties, maintenance rules and approval binding."],
        ["Inventory", "Immutable discovery snapshots, stable identities, deterministic reconciliation, topology and history."],
        ["Patch Catalog", "Approved metadata, digests, applicability inputs, baseline and exception management."],
        ["Workflow", "Persisted DAG, scheduling, locks, fencing, retries, reconciliation and compensation plans."],
        ["Managed Agent", "Outbound communication, typed discovery/execution operations, local journal and evidence."],
        ["ArtifactStore", "Logical artifact IDs over local disk, NFS or optional S3-compatible adapters."],
        ["Audit and Evidence", "Append-only decisions and immutable evidence references with checksums and retention."],
    ]
    add_table(doc, ["Component", "Responsibility"], components, [2300, 7060], font_size=9.3)

    doc.add_heading("4. OEM-style dynamic discovery", level=1)
    add_body(doc, "A dedicated OS account runs the managed agent on each host. The agent uses fixed read-only collectors and narrowly scoped privilege elevation where Oracle ownership requires execution as oracle, grid or root. It reports observations; it does not directly edit the central inventory.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(discovery_img), width=Inches(6.42))
    set_picture_alt_text(shape, "Discovery and reconciliation loop", "Five stages: collect, create immutable snapshot, reconcile stable identity, project current inventory, and decide compliance/readiness.")
    cap = doc.add_paragraph("Figure 2. Evidence-backed discovery and reconciliation loop", style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("4.1 Discovery sources", level=2)
    sources = [
        ["Host", "Machine identity, hostnames, interfaces, OS/kernel, architecture, mounts, capacity and clock health."],
        ["Oracle installation", "oraInst location, central/local inventory, oratab, process executable paths and service definitions."],
        ["Oracle home", "Canonical path, owner, version, OPatch version, installed patches and inventory attachment."],
        ["Database", "DBID, DB unique name, version, instance, CDB/PDB structure, role, open mode and SQL patch registry."],
        ["Network/services", "Listeners, endpoints, registered services, startup/service configuration and runtime state."],
        ["Topology", "Data Guard relationships, GI cluster, RAC instances/services, ASM and placement relationships."],
    ]
    add_table(doc, ["Source area", "Collected facts"], sources, [1900, 7460], font_size=9.2)

    doc.add_heading("4.2 Observed state versus desired state", level=2)
    add_table(doc,
              ["State owner", "Examples", "Update rule"],
              [
                  ["Agent-observed", "Version, home path, runtime role, open mode, patches, capacity", "Updated only from evidence-backed discovery"],
                  ["Administrator", "Environment, criticality, owner, window, approved policy", "Never overwritten by reconciliation"],
                  ["Integration", "CMDB ID, backup evidence, change ticket", "Owned by the named integration"],
                  ["Computed", "Compliance, staleness, confidence, blast radius", "Recomputed from versioned rules"],
              ],
              [1700, 4240, 3420], font_size=9.0)

    doc.add_heading("5. Inventory identity and reconciliation", level=1)
    add_body(doc, "Names and placements change. Hostnames, IP addresses, SIDs and Data Guard roles must therefore remain attributes rather than primary identities. Ambiguous matches are quarantined as identity conflicts; they are not automatically merged.")
    identities = [
        ["Agent", "Enrollment UUID bound to workload identity/certificate"],
        ["Host", "Installation identity plus OS machine evidence; clone detection required"],
        ["Oracle home", "Host identity + canonical real path + inventory identity/fingerprint"],
        ["Database", "DBID when available, reinforced by DB unique name and creation evidence; never SID alone"],
        ["Instance", "Database identity + instance number; node placement is mutable"],
        ["Cluster/topology", "Cluster/configuration identity; database role and service placement are mutable"],
    ]
    add_table(doc, ["Resource", "Identity strategy"], identities, [1900, 7460], font_size=9.2)
    doc.add_heading("5.1 Absence and partial discovery", level=2)
    for item in [
        "A failed or partial scan preserves the previous inventory and marks affected data inaccessible or unknown.",
        "One missing observation becomes not_observed_this_run, not deleted.",
        "Repeated complete scans with full source coverage may advance a resource to suspected_removed.",
        "Only deterministic policy plus adequate evidence may advance it to removed_confirmed.",
        "Stopped but configured databases remain configured_not_running.",
    ]:
        add_bullet(doc, item, bullet_id)

    doc.add_heading("6. Patch lifecycle and execution safety", level=1)
    add_body(doc, "Patching is modeled as a reconciled state machine rather than a sequence of remote commands. Each phase defines preconditions, a completion predicate, timeout behavior, interruption policy, evidence requirements and recovery boundary. Exit code zero alone is not proof of success.")
    phases = [
        ["1", "Freeze and rediscover", "Verify target, topology, inventory and plan revisions are still current."],
        ["2", "Validate artifact", "Verify logical artifact, digest, platform, metadata, staging safety and OPatch requirements."],
        ["3", "Run prechecks", "Evaluate conflicts, space, inventory health, database/PDB state, backup and recovery evidence."],
        ["4", "Approve", "Bind decision to immutable plan hash, target set, topology, artifacts, policy and window."],
        ["5", "Quiesce", "Drain approved services and prove the intended blast radius is inactive."],
        ["6", "Patch binaries", "Run a fixed Oracle adapter as the approved OS owner; capture durable local evidence."],
        ["7", "Verify binaries", "Re-observe inventory and confirm the exact expected patch state."],
        ["8", "Start and SQL patch", "Start in the required mode, run the approved datapatch profile and account for every required container."],
        ["9", "Postcheck", "Compare before/after state, services, registry, invalid objects, logs and smoke tests."],
        ["10", "Close or recover", "Update compliance after success or enter an approved recovery/manual-intervention plan."],
    ]
    add_table(doc, ["#", "Phase", "Required outcome"], phases, [500, 2300, 6560], font_size=8.9)

    doc.add_heading("6.1 Interrupted-operation rules", level=2)
    for item in [
        "A lease timeout means outcome unknown; it does not mean the Oracle operation failed.",
        "The agent may finish only the already-authorized atomic phase according to that phase's offline policy.",
        "No subsequent disruptive phase starts while the control plane, approval, window or evidence service is unavailable.",
        "Before a destructive retry, the reconciler classifies actual state as not started, running, complete, partial or indeterminate.",
        "Partial or indeterminate state pauses for a defined recovery procedure or manual intervention.",
    ]:
        add_bullet(doc, item, bullet_id)

    doc.add_heading("6.2 Rollback and recovery", level=2)
    add_body(doc, "Rollback is an approved, precompiled compensation plan—not an improvised reversal. Binary rollback, SQL rollback, out-of-place switchback, restore points, RMAN recovery, storage snapshots and Data Guard recovery provide different guarantees. The platform must never claim full recovery until binary, SQL, database and service states have been independently verified.")
    add_callout(doc, "Mandatory stop", "Central inventory inconsistency, ambiguous partial patch state, unclassified datapatch error, missing rollback evidence, unexpected topology change, or an unsupported README requirement must result in manual intervention rather than optimistic automation.", fill="FFF4F2", accent=RED)

    doc.add_heading("7. Security, trust and deployment", level=1)
    doc.add_heading("7.1 Agent trust boundary", level=2)
    for item in [
        "Outbound HTTPS with mutual TLS; no inbound SSH requirement.",
        "Unique agent workload identity, short-lived certificates, rotation, revocation and quarantine.",
        "Tasks bind the agent, target revision, operation/version, plan hash, policy snapshot, artifact digest, fencing token and expiry.",
        "Agent operations are compiled into the installed release; the server cannot upload executable code.",
        "A narrowly scoped privilege broker permits only registered actions as oracle, grid or root.",
        "Secrets are referenced, resolved just in time, and redacted before logs or evidence leave the host.",
    ]:
        add_bullet(doc, item, bullet_id)

    doc.add_heading("7.2 Roles and separation of duties", level=2)
    roles = [
        ["Platform administrator", "Configuration, identity integration and platform operation; no automatic patch authority."],
        ["Agent enrollment administrator", "Agent enrollment, certificate lifecycle and quarantine."],
        ["Patch catalog manager", "Ingests, verifies and publishes patch metadata/artifacts."],
        ["Planner/requester", "Selects targets and prepares an immutable change plan."],
        ["Approver", "Reviews blast radius, evidence, recovery and policy-bound plan."],
        ["Operator", "Runs an already approved plan and responds to controlled pauses."],
        ["Auditor/viewer", "Reads inventory, decisions, evidence and reports without mutation."],
    ]
    add_table(doc, ["Role", "Boundary"], roles, [2500, 6860], font_size=9.1)

    doc.add_heading("7.3 Deployment-neutral artifact storage", level=2)
    add_body(doc, "Business records contain an opaque artifact ID, digest, size, media type and lifecycle state. They never contain NFS paths, bucket names, provider URLs or SDK-specific types. The first adapters are local POSIX filesystem and shared POSIX/NFS. An S3-compatible adapter is optional and does not change patch, workflow or evidence records.")
    add_table(doc,
              ["Deployment profile", "Control plane", "Artifact storage", "Cloud dependency"],
              [
                  ["Developer/lab", "Single node + PostgreSQL", "Local filesystem", "None"],
                  ["On-premises production", "HA API/workers + PostgreSQL HA", "Shared filesystem/NFS", "None"],
                  ["Optional private/public cloud", "Same product services", "Optional S3-compatible adapter", "Optional"],
              ],
              [1700, 2950, 2850, 1860], font_size=8.8)

    doc.add_heading("8. Principal challenges and mitigations", level=1)
    challenges = [
        ["Inconsistent discovery sources", "Combine runtime, configuration, inventory and SQL evidence; retain provenance and contradictions."],
        ["Multiple OS owners", "Dedicated agent plus narrow typed privilege rules; discovery privileges remain separate from patch privileges."],
        ["Unstable names and roles", "Stable identity model; mutable hostnames, SIDs, role and node placement remain attributes."],
        ["Partial scans", "Coverage metadata and immutable snapshots; never infer deletion from failure or inaccessibility."],
        ["Shared Oracle homes", "Calculate blast radius and lock the home before approval or execution."],
        ["Patch applicability", "Use approved metadata and Oracle tools; unknown output or unsupported profile fails closed."],
        ["Lost agent/control connection", "Local durable supervisor plus postcondition reconciliation; no blind destructive retry."],
        ["Rollback uncertainty", "Precompiled recovery plans, retained evidence and lab drills; explicit manual-intervention states."],
        ["Data Guard/RAC complexity", "Separate topology adapters and release gates after standalone behavior is proven."],
        ["Powerful credentials", "mTLS, short-lived secret retrieval, least privilege, redaction and separation of duties."],
        ["Air-gapped estates", "Offline bundles, filesystem/NFS storage and no public runtime service requirement."],
        ["Testing real failures", "Simulator fault injection plus disposable standalone, Data Guard and RAC labs."],
    ]
    add_table(doc, ["Challenge", "Architectural response"], challenges, [2700, 6660], font_size=8.8)

    doc.add_heading("9. Accuracy and quality gates", level=1)
    gates = [
        ["Q0 Contract", "Canonical identities, ownership, schemas and compatibility fixtures pass."],
        ["Q1 Read-only safety", "No arbitrary shell/SQL path; privilege, validation, timeout, limit and redaction tests pass."],
        ["Q2 Discovery accuracy", "At least 99% entity/relationship correctness against verified standalone lab ground truth; no scan duplicates."],
        ["Q3 Degraded operation", "Partial permissions, unavailable sources and interrupted uploads preserve inventory and expose uncertainty."],
        ["Q4 Workflow simulation", "Failure injection at every node proves restart, duplicate delivery, lease, window, approval and rollback behavior."],
        ["Q5 Standalone lab", "Apply, conflict, no-space, agent loss, OPatch/datapatch failure and recovery scenarios retain complete evidence."],
        ["Q6 Topology", "Data Guard role change and RAC service movement preserve identity and apply only verified topology logic."],
        ["Q7 Pilot/production", "Threat model, recovery drill, DBA review, security review, runbooks, SLOs and abort criteria are approved."],
    ]
    add_table(doc, ["Gate", "Exit condition"], gates, [1900, 7460], font_size=9.1)

    doc.add_heading("10. Workstreams and parallel delivery waves", level=1)
    workstreams = [
        ["Contracts", "Canonical models, schemas, lifecycle and compatibility"],
        ["Agent", "Enrollment, identity, heartbeat, leasing, safe runner and upgrades"],
        ["Discovery", "Host, home, DB, listener, Data Guard and GI/RAC collectors"],
        ["Inventory", "Snapshots, identity, reconciliation, topology and history"],
        ["Artifact/compliance", "Storage abstraction, patch catalog, baseline and compliance"],
        ["Security", "RBAC, policy, approvals, secrets, audit and threat model"],
        ["Workflow", "Plan DAG, scheduling, locks, retries, reconciliation and compensation"],
        ["Readiness", "Fixed prechecks, recovery evidence and explainable decision"],
        ["Execution", "Staging, lifecycle, OPatch, datapatch, verification and recovery"],
        ["Experience", "API, UI, reporting, notifications and enterprise integrations"],
        ["Quality/platform", "Simulator, Oracle labs, CI, deployment, observability and operations"],
    ]
    add_table(doc, ["Workstream", "Primary ownership"], workstreams, [2300, 7060], font_size=9.2)

    doc.add_heading("10.1 Parallel waves", level=2)
    waves = [
        ["0", "Freeze system language", "Canonical IDs, schemas, result/event contracts, compatibility policy and initial fixtures."],
        ["1", "Parallel foundations", "Agent transport, inventory persistence, filesystem artifacts, security/audit, workflow kernel and test platform."],
        ["2", "Trustworthy inventory", "Standalone discovery, identity reconciliation, conflicts, history and accuracy measurement."],
        ["3", "Readiness product", "Patch metadata, compliance, plan compiler, policy, approvals, locks and read-only prechecks."],
        ["4", "Fault-injected simulator", "End-to-end lifecycle with duplicate, restart, expiry, storage and approval-drift failures."],
        ["5", "Standalone Oracle lab", "Controlled RU execution and recovery behind a disabled-by-default production gate."],
        ["6", "Pilot and operations", "UI, integrations, on-prem/air-gapped packaging, recovery drills and non-production canary."],
        ["7", "Topology expansion", "OJVM, Data Guard, then GI/RAC through separately tested profiles."],
    ]
    add_table(doc, ["Wave", "Objective", "Exit focus"], waves, [700, 2500, 6160], font_size=8.9)
    add_callout(doc, "Critical path", "CON-01 identities → versioned schemas → read-only agent → immutable snapshots → reconciliation → compliance/prechecks → workflow simulator → standalone lab execution → recovery drills → pilot.")

    doc.add_heading("11. First implementation package", level=1)
    doc.add_heading("CON-01 - Canonical resource identities", level=2)
    add_body(doc, "Identity is the first implementation package because every later component must agree on what a host, Oracle home, database, instance, listener, cluster and topology configuration actually is. Incorrect identity causes duplicate inventory, unsafe merges, wrong blast-radius calculations and potentially the wrong patch target.")
    doc.add_heading("11.1 Deliverables", level=2)
    first_deliverables = [
        "A normative identity specification for agent, host, Oracle home, database, instance, listener, cluster, ASM and Data Guard configuration.",
        "Immutable internal ID formats and external/natural-key evidence used for matching.",
        "Rules distinguishing identity from mutable attributes such as hostname, IP, SID, role, node placement and service endpoint.",
        "Clone detection, ambiguous-match quarantine, manual merge/split and audit semantics.",
        "Valid, invalid and collision-oriented JSON fixtures for standalone, cloned-host, home-relocation, SID-reuse, Data Guard switchover and RAC movement cases.",
        "An architecture decision record approved by discovery, inventory, agent, workflow, DBA and security owners.",
    ]
    for item in first_deliverables:
        add_bullet(doc, item, bullet_id)

    doc.add_heading("11.2 Acceptance criteria", level=2)
    criteria = [
        "No resource uses only hostname, IP address or SID as its identity.",
        "The same database after restart, switchover or instance movement keeps its identity.",
        "A cloned VM or Oracle home is detected and quarantined rather than silently merged.",
        "Identity decisions expose reason codes and match evidence.",
        "Ambiguous inputs produce identity_conflict, not a guessed match.",
        "All downstream teams can reference the same versioned identifiers without importing inventory internals.",
    ]
    for item in criteria:
        add_bullet(doc, item, bullet_id)

    doc.add_heading("11.3 Immediate parallel preparation", level=2)
    add_body(doc, "While CON-01 is being reviewed, teams may prepare—but not finalize—FND-01 repository boundaries, SEC-03 threat scenarios, QA-01 fixture format, and the design skeletons for agent, snapshot, artifact and workflow contracts. They must not invent competing identifiers.")

    doc.add_heading("12. Governance and definition of done", level=1)
    for item in [
        "One accountable owner and reviewer set per package.",
        "Interface and acceptance tests are written before implementation is considered complete.",
        "Parser changes include sanitized golden Oracle output and unknown-output failure tests.",
        "High-risk execution changes require independent DBA and security/platform review.",
        "Every included destructive phase has an independently observable completion predicate.",
        "Apply and recovery are demonstrated repeatedly from clean lab images.",
        "All transitions can be reconstructed from audit and immutable evidence.",
        "Production execution is explicitly enabled by policy; installing the software never enables patching automatically.",
    ]:
        add_bullet(doc, item, bullet_id)

    doc.add_heading("13. Decisions to validate", level=1)
    decisions = [
        ["Identity provider", "Enterprise IdP/protocol and service-account lifecycle"],
        ["Secret provider", "Enterprise vault or controlled local provider for first deployment"],
        ["Agent OS account", "Standard account name, installation path, ownership and service manager"],
        ["Privilege model", "Approved sudo/polkit mechanism and oracle/grid/root operation boundary"],
        ["Database connectivity", "Local OS authentication versus managed credential/wallet profiles"],
        ["Artifact deployment", "Local filesystem or shared NFS path, retention, capacity and backup"],
        ["Backup evidence", "RMAN catalog, database views, enterprise backup API or approved combination"],
        ["Service control", "Application drain/start integration boundary per environment"],
        ["Change management", "Whether and when an external change system becomes authoritative"],
        ["Initial Oracle lab", "Exact 19c RU baseline, CDB/non-CDB, shared-home and failure scenarios"],
        ["Pilot boundary", "First non-production hosts, abort criteria and required sign-offs"],
    ]
    add_table(doc, ["Decision", "Required clarification"], decisions, [2600, 6760], font_size=9.0)

    # Appendix A - full work package register parsed from source.
    doc.add_page_break()
    doc.add_heading("Appendix A. Complete work-package register", level=1)
    add_body(doc, "The following register is the authoritative assignment index at this architecture baseline. Package dependencies must be satisfied before implementation is merged. Acceptance evidence is mandatory; code completion alone does not close a package.")
    tracks = parse_work_packages(ROOT / "Oracle_Patching_Utility_Blueprint.md")
    package_count = 0
    for track in tracks:
        if not track["rows"]:
            continue
        doc.add_heading(track["title"], level=2)
        add_table(doc, ["ID", "Package", "Depends on", "Acceptance evidence"], track["rows"],
                  [1050, 3400, 1900, 3010], font_size=8.15, alternate=True)
        package_count += len(track["rows"])
    add_callout(doc, "Register total", f"{package_count} unique work packages across the complete greenfield program.")

    doc.add_heading("Appendix B. Core terminology", level=1)
    glossary = [
        ["Observation", "One evidence-backed fact produced by a versioned collector."],
        ["Discovery snapshot", "Immutable set of observations, source coverage and collection errors from one run."],
        ["Inventory projection", "Current resource/topology state computed by deterministic reconciliation."],
        ["Desired state", "Administrator-approved baseline, policy and business metadata."],
        ["Artifact", "Immutable patch archive or evidence content addressed logically by ID and digest."],
        ["Change request", "Business authorization, scope, maintenance window and required approvals."],
        ["Patch plan", "Immutable compiled graph bound to target/topology revisions, policy and artifact digests."],
        ["Workflow run", "One persisted execution of an approved patch plan."],
        ["Task", "One typed operation addressed to an agent and target under a workflow node."],
        ["Lease", "Time-bounded ownership of one task delivery; expiry means outcome unknown."],
        ["Attempt", "One supervised execution attempt with a stable idempotency key and evidence."],
        ["Reconciliation", "Observation of actual state to determine identity, current inventory or execution outcome."],
        ["Completion predicate", "Machine-verifiable observed condition required to declare a phase successful."],
        ["Compensation plan", "Precompiled, approved recovery graph; not an improvised reversal."],
        ["Manual intervention", "Safe terminal/pause state requiring an operator or DBA procedure."],
    ]
    add_table(doc, ["Term", "Meaning"], glossary, [2300, 7060], font_size=9.1)

    # Final note.
    doc.add_heading("Implementation authorization point", level=1)
    add_callout(doc, "Ready to begin", "Start CON-01 canonical identities. Do not implement Oracle mutation adapters until the contract, discovery, reconciliation, simulator and quality-gate prerequisites in this blueprint have passed.", fill="ECF8F0", accent=GREEN)
    doc.add_heading("Immediate kickoff checklist", level=2)
    add_table(doc,
              ["Action", "Required output"],
              [
                  ["Appoint CON-01 owner and reviewers", "Named owners from architecture, discovery, inventory, agent, DBA and security"],
                  ["Run the identity workshop", "Agreed identity versus mutable-attribute rules for every canonical resource"],
                  ["Create the collision fixture set", "Clone, rename, SID reuse, home relocation, role change and node movement cases"],
                  ["Approve the identity ADR", "Versioned decision and acceptance record usable by every downstream package"],
              ],
              [3300, 6060], font_size=9.2)
    add_body(doc, "Allowed parallel preparation: FND-01 repository boundaries, SEC-03 threat scenarios, QA-01 fixture format and draft contract skeletons may begin while CON-01 is reviewed. They must not finalize competing identifiers.")

    # Core properties and save.
    props = doc.core_properties
    props.title = "Oracle Patching Utility - Architecture and Delivery Blueprint"
    props.subject = "Greenfield OEM-style Oracle discovery and patching platform"
    props.author = "Oracle Patching Utility Program"
    props.keywords = "Oracle, patching, discovery, agent, architecture, work packages"
    props.comments = "Architecture baseline for implementation"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
